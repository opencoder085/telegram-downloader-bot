import os
import re
import difflib
import json
import ssl
import shutil
import asyncio
import tempfile
import threading
import time
import struct
import urllib.parse
import urllib.request
import calendar
from datetime import datetime, timedelta
from http.server import HTTPServer, BaseHTTPRequestHandler
import httpx
from PIL import Image, ImageDraw, ImageFont
import pypdf
from pypdf import PdfWriter, PdfReader
import docx
import openpyxl
import reportlab
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import pytesseract
from telegram import (
    Update,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    ReplyKeyboardMarkup,
    KeyboardButton,
    BotCommand,
    BotCommandScopeChat,
)
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)
import yt_dlp

# =====================================================================
# RENDER 7/24 SAĞLIK KONTROLÜ VE SELF-PINGER
# =====================================================================
class HealthCheckHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-type", "text/plain; charset=utf-8")
        self.end_headers()
        self.wfile.write(b"NUN BOT 7/24 AKTIF!")

    def do_HEAD(self):
        self.send_response(200)
        self.end_headers()

    def log_message(self, format, *args):
        return

def run_health_server():
    port = int(os.environ.get("PORT", 8080))
    server = HTTPServer(('0.0.0.0', port), HealthCheckHandler)
    server.serve_forever()

def run_keep_alive_pinger():
    target_url = os.environ.get("RENDER_EXTERNAL_URL") or os.environ.get("APP_URL")
    if not target_url:
        return
    if not target_url.startswith("http"):
        target_url = f"https://{target_url}"

    while True:
        time.sleep(540)
        try:
            req = urllib.request.Request(target_url, headers={'User-Agent': 'NunBot-KeepAlive/1.0'})
            with urllib.request.urlopen(req, timeout=20) as resp:
                pass
        except Exception:
            pass

# =====================================================================
# VERİ TABANI YÖNETİMİ (DİL, SINAVLAR, HATIRLATICILAR)
# =====================================================================
LANG_FILE = "user_langs.json"
EXAMS_FILE = "user_exams.json"
REMINDERS_FILE = "user_reminders.json"
USER_LANGS = {}
USER_EXAMS = {}
USER_REMINDERS = {}

def load_databases():
    global USER_LANGS, USER_EXAMS, USER_REMINDERS
    for fname, var_ref in [(LANG_FILE, USER_LANGS), (EXAMS_FILE, USER_EXAMS), (REMINDERS_FILE, USER_REMINDERS)]:
        if os.path.exists(fname):
            try:
                with open(fname, "r", encoding="utf-8") as f:
                    var_ref.update(json.load(f))
            except Exception:
                pass

def save_json(fname, data):
    try:
        with open(fname, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
    except Exception:
        pass

def save_user_lang(user_id, lang_code: str):
    USER_LANGS[str(user_id)] = lang_code
    save_json(LANG_FILE, USER_LANGS)

def get_user_lang(user_id, context: ContextTypes.DEFAULT_TYPE = None) -> str:
    uid_str = str(user_id)
    if uid_str in USER_LANGS and USER_LANGS[uid_str] in TEXTS:
        return USER_LANGS[uid_str]
    if context and context.user_data and 'lang' in context.user_data:
        l = context.user_data['lang']
        if isinstance(l, str) and l in TEXTS:
            return l
    return 'uz'

def add_user_exam(user_id, title: str, date_str: str) -> str:
    uid_str = str(user_id)
    if uid_str not in USER_EXAMS:
        USER_EXAMS[uid_str] = []
    exam_id = str(int(time.time() * 1000))[-6:]
    USER_EXAMS[uid_str].append({'id': exam_id, 'title': title, 'date': date_str})
    save_json(EXAMS_FILE, USER_EXAMS)
    return exam_id

def delete_user_exam(user_id, exam_id: str) -> bool:
    uid_str = str(user_id)
    if uid_str in USER_EXAMS:
        orig = len(USER_EXAMS[uid_str])
        USER_EXAMS[uid_str] = [e for e in USER_EXAMS[uid_str] if e.get('id') != exam_id]
        if len(USER_EXAMS[uid_str]) < orig:
            save_json(EXAMS_FILE, USER_EXAMS)
            return True
    return False

def get_user_exams(user_id) -> list:
    return USER_EXAMS.get(str(user_id), [])

def add_user_reminder(user_id, text: str, target_dt_str: str) -> str:
    uid_str = str(user_id)
    if uid_str not in USER_REMINDERS:
        USER_REMINDERS[uid_str] = []
    rem_id = str(int(time.time() * 1000))[-6:]
    USER_REMINDERS[uid_str].append({'id': rem_id, 'text': text, 'time': target_dt_str})
    save_json(REMINDERS_FILE, USER_REMINDERS)
    return rem_id

def get_user_reminders(user_id) -> list:
    return USER_REMINDERS.get(str(user_id), [])

def delete_user_reminder(user_id, rem_id: str):
    uid_str = str(user_id)
    if uid_str in USER_REMINDERS:
        USER_REMINDERS[uid_str] = [r for r in USER_REMINDERS[uid_str] if r.get('id') != rem_id]
        save_json(REMINDERS_FILE, USER_REMINDERS)

def cleanup_user_temp_files(context, user_id):
    if not context or not context.user_data:
        return
    for f in context.user_data.get('merge_files', []):
        try:
            if os.path.exists(f): os.remove(f)
        except Exception: pass
    context.user_data.pop('merge_files', None)

    sp = context.user_data.get('split_file_path')
    if sp:
        try:
            if os.path.exists(sp): os.remove(sp)
        except Exception: pass
    context.user_data.pop('split_file_path', None)
    context.user_data.pop('split_max_pages', None)

    dp = context.user_data.get('direct_file_path')
    if dp:
        try:
            if os.path.exists(dp): os.remove(dp)
        except Exception: pass
    context.user_data.pop('direct_file_path', None)

    context.user_data.pop('exam_draft_title', None)
    context.user_data.pop('exam_draft_dt', None)
    context.user_data['mode'] = 'auto'

# =====================================================================
# PDF & BELGE DÖNÜŞTÜRME MOTORLARI
# =====================================================================
def docx_to_pdf(input_docx: str, output_pdf: str) -> bool:
    try:
        doc = docx.Document(input_docx)
        styles = getSampleStyleSheet()
        story = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                story.append(Spacer(1, 8))
                continue
            if p.style.name.startswith('Heading'):
                story.append(Paragraph(f"<b>{txt}</b>", styles['Heading2']))
                story.append(Spacer(1, 6))
            else:
                story.append(Paragraph(txt, styles['Normal']))
                story.append(Spacer(1, 4))
        pdf_doc = SimpleDocTemplate(output_pdf, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        pdf_doc.build(story)
        return True
    except Exception as e:
        print(f"docx_to_pdf hatasi: {e}")
        return False

def xlsx_to_pdf(input_xlsx: str, output_pdf: str) -> bool:
    try:
        wb = openpyxl.load_workbook(input_xlsx, data_only=True)
        ws = wb.active
        table_data = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                table_data.append([str(c) if c is not None else "" for c in row])
        if not table_data:
            return False
        pdf_doc = SimpleDocTemplate(output_pdf, pagesize=A4)
        story = [Paragraph("<b>Excel Jadvali</b>", getSampleStyleSheet()['Heading2']), Spacer(1, 10)]
        t = Table(table_data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#222222')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('GRID', (0,0), (-1,-1), 1, colors.HexColor('#666666')),
            ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ]))
        story.append(t)
        pdf_doc.build(story)
        return True
    except Exception as e:
        print(f"xlsx_to_pdf hatasi: {e}")
        return False

def txt_to_pdf(input_txt: str, output_pdf: str) -> bool:
    try:
        with open(input_txt, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        doc = SimpleDocTemplate(output_pdf, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        story = [Paragraph("<b>Matn Hujjati</b>", styles['Heading2']), Spacer(1, 10)]
        for line in content.split("\n"):
            line_clean = line.strip()
            if line_clean:
                story.append(Paragraph(line_clean, styles['Normal']))
                story.append(Spacer(1, 4))
        doc.build(story)
        return True
    except Exception as e:
        print(f"txt_to_pdf hatasi: {e}")
        return False

def pdf_to_docx(input_pdf: str, output_docx: str) -> bool:
    try:
        reader = PdfReader(input_pdf)
        doc = docx.Document()
        doc.add_heading("PDF'dan Oʻgirilgan Hujjat", level=1)
        has_text = False
        for page in reader.pages:
            txt = page.extract_text()
            if txt:
                has_text = True
                for p in txt.split("\n"):
                    if p.strip():
                        doc.add_paragraph(p.strip())
        if not has_text:
            return False
        doc.save(output_docx)
        return True
    except Exception as e:
        print(f"pdf_to_docx hatasi: {e}")
        return False

def merge_pdfs(pdf_paths: list, output_pdf: str) -> bool:
    try:
        writer = PdfWriter()
        for p in pdf_paths:
            writer.append(p)
        with open(output_pdf, "wb") as f:
            writer.write(f)
        return True
    except Exception as e:
        print(f"merge_pdfs hatasi: {e}")
        return False

def parse_page_ranges(range_str: str, max_pages: int):
    selected = set()
    parts = range_str.replace(" ", "").split(",")
    for p in parts:
        if "-" in p:
            start_str, _, end_str = p.partition("-")
            if start_str.isdigit() and end_str.isdigit():
                start = max(1, int(start_str))
                end = min(max_pages, int(end_str))
                for pg in range(start, end + 1):
                    selected.add(pg - 1)
        elif p.isdigit():
            pg = int(p)
            if 1 <= pg <= max_pages:
                selected.add(pg - 1)
    return sorted(list(selected))

def split_pdf(input_pdf: str, range_str: str, output_pdf: str) -> bool:
    try:
        reader = PdfReader(input_pdf)
        max_p = len(reader.pages)
        pages_to_extract = parse_page_ranges(range_str, max_p)
        if not pages_to_extract:
            return False
        writer = PdfWriter()
        for idx in pages_to_extract:
            writer.add_page(reader.pages[idx])
        with open(output_pdf, "wb") as f:
            writer.write(f)
        return True
    except Exception as e:
        print(f"split_pdf hatasi: {e}")
        return False

# =====================================================================
# TESSERACT OCR MOTORU
# =====================================================================
def extract_text_from_image(image_path: str) -> str:
    try:
        img = Image.open(image_path)
        avail = []
        try:
            avail = pytesseract.get_languages()
        except Exception:
            avail = ['eng']
        langs = [l for l in ['tur', 'uzb', 'rus', 'eng'] if l in avail]
        lang_arg = "+".join(langs) if langs else "eng"
        text = pytesseract.image_to_string(img, lang=lang_arg)
        return text.strip()
    except Exception as e:
        print(f"OCR hatasi: {e}")
        return ""

# =====================================================================
# HAFTALIK DERS PROGRAMI GÖRSELİ (PILLOW / 1080x1920)
# =====================================================================
def parse_schedule_text(text: str):
    day_keywords = {
        "dushanba": "DUSHANBA", "pazartesi": "PAZARTESİ", "monday": "MONDAY", "понедельник": "ПОНЕДЕЛЬНИК",
        "seshanba": "SESHANBA", "sali": "SALI", "tuesday": "TUESDAY", "вторник": "ВТОРНИК",
        "chorshanba": "CHORSHANBA", "carsamba": "ÇARŞAMBA", "wednesday": "WEDNESDAY", "среда": "СРЕДА",
        "payshanba": "PAYSHANBA", "persembe": "PERŞEMBE", "thursday": "THURSDAY", "четверг": "ЧЕТВЕРГ",
        "juma": "JUMA", "cuma": "CUMA", "friday": "FRIDAY", "пятница": "ПЯТНИЦA",
        "shanba": "SHANBA", "cumartesi": "CUMARTESİ", "saturday": "SATURDAY", "суббота": "СУББОТА",
        "yakshanba": "YAKSHANBA", "pazar": "PAZAR", "sunday": "SUNDAY", "воскресенье": "ВОСКРЕСЕНЬЕ",
    }
    current_day = "DUSHANBA"
    schedule = {}
    lines = text.strip().split("\n")
    for line in lines:
        l_c = line.strip()
        if not l_c:
            continue
        first_word = l_c.lower().replace(":", "").replace("-", "").split()[0]
        if first_word in day_keywords:
            current_day = day_keywords[first_word]
            if current_day not in schedule:
                schedule[current_day] = []
            rest = l_c[len(first_word):].lstrip(":- ").strip()
            if rest:
                schedule[current_day].append(rest)
        else:
            if current_day not in schedule:
                schedule[current_day] = []
            schedule[current_day].append(l_c)
    return schedule

def generate_schedule_wallpaper(schedule_data: dict, output_path: str, title: str = "DERS PROGRAMI"):
    width, height = 1080, 1920
    img = Image.new("RGB", (width, height), color=(10, 10, 10))
    draw = ImageDraw.Draw(img)

    try:
        font_large = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 46)
        font_med = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 32)
        font_sub = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 24)
        font_item = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 26)
    except Exception:
        font_large = font_med = font_sub = font_item = ImageFont.load_default()

    for x in range(0, width, 80):
        draw.line([(x, 0), (x, height)], fill=(18, 18, 18), width=1)
    for y in range(0, height, 80):
        draw.line([(0, y), (width, y)], fill=(18, 18, 18), width=1)

    draw.rectangle([(60, 80), (width - 60, 220)], fill=(18, 18, 18), outline=(60, 60, 60), width=2)
    draw.text((100, 105), "NUN PROJECT // ACADEMIC", fill=(160, 160, 160), font=font_sub)
    draw.text((100, 140), title.upper(), fill=(255, 255, 255), font=font_large)

    cur_y = 260
    for day, items in schedule_data.items():
        if cur_y > height - 200:
            break
        day_box_height = 55 + (len(items) * 45)
        draw.rectangle([(60, cur_y), (width - 60, cur_y + day_box_height)], fill=(15, 15, 15), outline=(50, 50, 50), width=1)
        draw.rectangle([(60, cur_y), (width - 60, cur_y + 45)], fill=(30, 30, 30))
        draw.text((80, cur_y + 8), f"●  {day}", fill=(255, 255, 255), font=font_med)

        item_y = cur_y + 55
        for it in items:
            draw.text((90, item_y + 5), f"▫️  {it}", fill=(220, 220, 220), font=font_item)
            item_y += 42
        cur_y += day_box_height + 20

    draw.text((width // 2 - 180, height - 70), "NUN PROJECT • LOCKSCREEN TIMETABLE", fill=(100, 100, 100), font=font_sub)
    img.save(output_path, quality=95)
    return True

# =====================================================================
# POMODORO & HATIRLATICI MOTORU
# =====================================================================
async def pomodoro_timer_task(bot, chat_id: int, duration_mins: int, is_break: bool, user_lang: str):
    await asyncio.sleep(duration_mins * 60)
    if is_break:
        msg = "🔔 *TANAFFUS TUGADI!* ☕\n\nYangi ish seansiga tayyormisiz?" if user_lang == 'uz' else "🔔 *MOLA BİTTİ!* ☕\n\nYeni çalışma seansına hazır mısınız?"
        btn_text = "🍅 25 Dk Pomodoro"
        cb = "pomo_25"
    else:
        msg = f"🎉 *POMODORO TUGADI!* 🍅\n\n`{duration_mins}` daqiqalik dars yakunlandi!\nEndi 5 daqiqa dam oling." if user_lang == 'uz' else f"🎉 *SÜRE BİTTİ!* 🍅\n\n`{duration_mins}` dakikalık seans tamamlandı!\nŞimdi 5 dakika mola vakti."
        btn_text = "☕ 5 Dk Mola"
        cb = "pomo_5"

    kb = InlineKeyboardMarkup([[InlineKeyboardButton(btn_text, callback_data=cb)]])
    try:
        await bot.send_message(chat_id=chat_id, text=msg, parse_mode="Markdown", reply_markup=kb)
    except Exception:
        pass

async def reminders_worker(app):
    while True:
        await asyncio.sleep(30)
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        due = []
        for uid, items in list(USER_REMINDERS.items()):
            for it in list(items):
                if it.get("time") <= now_str:
                    due.append((uid, it))
                    items.remove(it)
        if due:
            save_json(REMINDERS_FILE, USER_REMINDERS)
            for uid, item in due:
                try:
                    await app.bot.send_message(
                        chat_id=int(uid),
                        text=f"🔔 *NUN PROJECT // ESLATMA*\n\n📌 *Vazifa:* {item.get('text')}\n⏰ Belgilangan vaqt yetib keldi!",
                        parse_mode="Markdown"
                    )
                except Exception:
                    pass

# =====================================================================
# ÖZBEKÇE KİRİL <-> LATİN ÇEVİRİ MOTORU
# =====================================================================
APOSTROPHES = set(["'", "\u2019", "\u2018", "`", "\u02bb", "\u02bc"])
VOWELS_CYR = set("аоуиэеёюяўАОУИЭЕЁЮЯЎ")

MAP_CYR_TO_LAT = {
    '\u0410': 'A', '\u0430': 'a', '\u0411': 'B', '\u0431': 'b', '\u0412': 'V', '\u0432': 'v',
    '\u0413': 'G', '\u0433': 'g', '\u0414': 'D', '\u0434': 'd', '\u0416': 'J', '\u0436': 'j',
    '\u0417': 'Z', '\u0437': 'z', '\u0418': 'I', '\u0438': 'i', '\u0419': 'Y', '\u0439': 'y',
    '\u041a': 'K', '\u043a': 'k', '\u049a': 'Q', '\u049b': 'q', '\u041b': 'L', '\u043b': 'l',
    '\u041c': 'M', '\u043c': 'm', '\u041d': 'N', '\u043d': 'n', '\u041e': 'O', '\u043e': 'o',
    '\u041f': 'P', '\u043f': 'p', '\u0420': 'R', '\u0440': 'r', '\u0421': 'S', '\u0441': 's',
    '\u0422': 'T', '\u0442': 't', '\u0423': 'U', '\u0443': 'u', '\u0424': 'F', '\u0444': 'f',
    '\u0425': 'X', '\u0445': 'x', '\u04b2': 'H', '\u04b3': 'h', '\u042d': 'E', '\u044d': 'e',
}

MAP_LAT_TO_CYR = {
    'A': '\u0410', 'a': '\u0430', 'B': '\u0411', 'b': '\u0431', 'V': '\u0412', 'v': '\u0432',
    'G': '\u0413', 'g': '\u0433', 'D': '\u0414', 'd': '\u0434', 'J': '\u0416', 'j': '\u0436',
    'Z': '\u0417', 'z': '\u0437', 'I': '\u0418', 'i': '\u0438', 'Y': '\u0419', 'y': '\u0439',
    'K': '\u041a', 'k': '\u043a', 'Q': '\u049a', 'q': '\u049b', 'L': '\u041b', 'l': '\u043b',
    'M': '\u041c', 'm': '\u043c', 'N': '\u041d', 'n': '\u043d', 'O': '\u041e', 'o': '\u043e',
    'P': '\u041f', 'p': '\u043f', 'R': '\u0420', 'r': '\u0440', 'S': '\u0421', 's': '\u0441',
    'T': '\u0422', 't': '\u0442', 'U': '\u0423', 'u': '\u0443', 'F': '\u0424', 'f': '\u0444',
    'X': '\u0425', 'x': '\u0445', 'H': '\u04b2', 'h': '\u04b3',
}

def cyrillic_to_latin(text: str) -> str:
    res = []
    i, n = 0, len(text)
    while i < n:
        c = text[i]
        prev = text[i-1] if i > 0 else " "
        nxt = text[i+1] if i+1 < n else ""
        if c in ('С', 'с') and nxt in ('Ҳ', 'ҳ'):
            res.append("S'H" if c.isupper() and nxt.isupper() else ("S'h" if c.isupper() else "s'h"))
            i += 2; continue
        if c in ('Е', 'е'):
            if i == 0 or not prev.isalpha() or prev in VOWELS_CYR or prev in 'ъЪьЬ':
                res.append("YE" if c.isupper() and nxt.isupper() else ("Ye" if c.isupper() else "ye"))
            else:
                res.append("E" if c.isupper() else "e")
            i += 1; continue
        if c in ('Ё', 'ё'): res.append("Yo" if c.isupper() else "yo"); i += 1; continue
        if c in ('Ю', 'ю'): res.append("Yu" if c.isupper() else "yu"); i += 1; continue
        if c in ('Я', 'я'): res.append("Ya" if c.isupper() else "ya"); i += 1; continue
        if c in ('Ч', 'ч'): res.append("Ch" if c.isupper() else "ch"); i += 1; continue
        if c in ('Ш', 'ш', 'Щ', 'щ'): res.append("Sh" if c.isupper() else "sh"); i += 1; continue
        if c in ('Ц', 'ц'): res.append("Ts" if c.isupper() else "ts"); i += 1; continue
        if c == 'Ў': res.append("Oʻ"); i += 1; continue
        if c == 'ў': res.append("oʻ"); i += 1; continue
        if c == 'Ғ': res.append("Gʻ"); i += 1; continue
        if c == 'ғ': res.append("gʻ"); i += 1; continue
        if c in ('Ъ', 'ъ'): res.append("'"); i += 1; continue
        if c in ('Ь', 'ь'): i += 1; continue
        res.append(MAP_CYR_TO_LAT.get(c, c))
        i += 1
    return "".join(res)

def latin_to_cyrillic(text: str) -> str:
    res = []
    i, n = 0, len(text)
    while i < n:
        c = text[i]
        nxt = text[i+1] if i+1 < n else ""
        nxt2 = text[i+2] if i+2 < n else ""
        prev = text[i-1] if i > 0 else " "
        if c in ('s', 'S') and nxt in APOSTROPHES and nxt2 in ('h', 'H'):
            res.append("СҲ" if nxt2.isupper() else "Сҳ" if c == 'S' else "сҳ"); i += 3; continue
        if c in ('o', 'O') and nxt and nxt in APOSTROPHES:
            res.append("Ў" if c == 'O' else "ў"); i += 2; continue
        if c in ('g', 'G') and nxt and nxt in APOSTROPHES:
            res.append("Ғ" if c == 'G' else "ғ"); i += 2; continue
        if c in ('s', 'S') and nxt in ('h', 'H'):
            res.append("Ш" if c.isupper() else "ш"); i += 2; continue
        if c in ('c', 'C') and nxt in ('h', 'H'):
            res.append("Ч" if c.isupper() else "ч"); i += 2; continue
        if c in ('t', 'T') and nxt in ('s', 'S'):
            res.append("Ц" if c.isupper() else "ц"); i += 2; continue
        if c in ('y', 'Y') and nxt in ('o', 'O'): res.append("Ё" if c.isupper() else "ё"); i += 2; continue
        if c in ('y', 'Y') and nxt in ('u', 'U'): res.append("Ю" if c.isupper() else "ю"); i += 2; continue
        if c in ('y', 'Y') and nxt in ('a', 'A'): res.append("Я" if c.isupper() else "я"); i += 2; continue
        if c in ('y', 'Y') and nxt in ('e', 'E'): res.append("Е" if c.isupper() else "е"); i += 2; continue
        if c in ('e', 'E'):
            res.append("Э" if (i == 0 or not prev.isalpha() or prev.lower() in 'aouie') else "Е" if c == 'E' else "е")
            i += 1; continue
        if c in APOSTROPHES:
            res.append("ъ" if prev.isalpha() else "'"); i += 1; continue
        res.append(MAP_LAT_TO_CYR.get(c, c))
        i += 1
    return "".join(res)

def is_mostly_cyrillic(text: str) -> bool:
    return len(re.findall(r'[\u0400-\u04FF]', text)) >= len(re.findall(r'[a-zA-Z]', text))

# =====================================================================
# NAMAZ VAKİTLERİ
# =====================================================================
UZ_REGIONS = {
    "toshkent": "toshkent", "samarqand": "samarqand-shahri", "buxoro": "buxoro-shahri",
    "andijon": "andijon-shahri", "namangan": "namangan-shahri", "fargona": "fargona-shahri",
    "qoqon": "qoqon-shahri", "urganch": "urganch-shahri", "nukus": "nukus-shahri",
    "qarshi": "qarshi-shahri", "navoiy": "navoiy-shahri", "termiz": "termiz-shahri",
}
GLOBAL_CITIES = {
    "istanbul": "Istanbul", "ankara": "Ankara", "izmir": "Izmir", "bursa": "Bursa",
    "moskva": "Moscow", "almaty": "Almaty", "makka": "Makkah", "madina": "Medina", "dubai": "Dubai"
}

def clean_time_str(val: str) -> str:
    m = re.search(r"\d{1,2}:\d{2}", str(val))
    return m.group(0) if m else "--:--"

async def fetch_prayer_times(city_input: str):
    c_norm = re.sub(r"['’`ʻʼ]", "", city_input.lower().strip()).replace('i̇', 'i').replace('ı', 'i')
    headers = {"User-Agent": "Mozilla/5.0"}

    async with httpx.AsyncClient(timeout=10.0, verify=False, follow_redirects=True) as client:
        slug = UZ_REGIONS.get(c_norm)
        if not slug:
            matches = difflib.get_close_matches(c_norm, list(UZ_REGIONS.keys()), n=1, cutoff=0.75)
            if matches: slug = UZ_REGIONS[matches[0]]
        if slug:
            try:
                resp = await client.get(f"https://namoz-vaqti.uz/index.php?format=json&region={slug}", headers=headers)
                if resp.status_code == 200:
                    j = resp.json()
                    t = j.get("today", {}).get("times", {})
                    meta = j.get("meta", {})
                    return {
                        "Fajr": t.get("bomdod"), "Sunrise": t.get("quyosh"), "Dhuhr": t.get("peshin"),
                        "Asr": t.get("asr"), "Maghrib": t.get("shom"), "Isha": t.get("xufton")
                    }, meta.get("region", {}).get("name", city_input.title()), meta.get("date", ""), "", "Oʻzbekiston Din ishlari qoʻmitasi"
            except Exception: pass

        target = GLOBAL_CITIES.get(c_norm, city_input)
        try:
            resp = await client.get(f"https://api.aladhan.com/v1/timingsByAddress?address={urllib.parse.quote(target)}", headers=headers)
            if resp.status_code == 200:
                data = resp.json().get("data", {})
                d = data.get("date", {})
                return data.get("timings", {}), target.title(), d.get("readable", ""), d.get("hijri", {}).get("date", ""), "AlAdhan API"
        except Exception: pass

    return None, None, None, None, None

def format_prayer_card(display_name: str, timings: dict, date_str: str, hijri_str: str, source: str, lang: str = 'uz'):
    t_f = clean_time_str(timings.get("Fajr"))
    t_s = clean_time_str(timings.get("Sunrise"))
    t_d = clean_time_str(timings.get("Dhuhr"))
    t_a = clean_time_str(timings.get("Asr"))
    t_m = clean_time_str(timings.get("Maghrib"))
    t_i = clean_time_str(timings.get("Isha"))
    lbls = ("BOMDOD", "QUYOSH", "PESHIN", "ASR", "SHOM", "XUFTON") if lang != 'tr' else ("İMSAK", "GÜNEŞ", "ÖĞLE", "İKİNDİ", "AKŞAM", "YATSI")
    return (
        f"*NUN PROJECT // NAMOZ VAQTLARI*\n"
        f"📍 *[ {display_name.upper()} ]*\n"
        f"📅 `{date_str}`\n\n"
        f"┌────────────────────────────┐\n"
        f"  ▫️ *{lbls[0]}:*    `{t_f}`\n"
        f"  ▫️ *{lbls}:*    `{t_s}`\n"
        f"  ▫️ *{lbls}:*    `{t_d}`\n"
        f"  ▫️ *{lbls}:*    `{t_a}`\n"
        f"  ▫️ *{lbls}:*    `{t_m}`\n"
        f"  ▫️ *{lbls}:*    `{t_i}`\n"
        f"└────────────────────────────┘\n"
        f"_{source}_"
    )

# =====================================================================
# CANLI İNTERAKTİF ZAMANLAYICI PANELİ & TAKVİM
# =====================================================================
def format_scheduler_card(title: str, dt: datetime, lang: str = 'uz') -> str:
    month_names = ["", "Yanvar", "Fevral", "Mart", "Aprel", "May", "Iyun", "Iyul", "Avgust", "Sentyabr", "Oktyabr", "Noyabr", "Dekabr"]
    clean_title = re.sub(r"[*_`\[\]]", "", title)
    return (
        f"*NUN PROJECT // VAQTNI SOZLASH*\n"
        f"📌 Fan / Imtihon: *{clean_title}*\n\n"
        f"┌──────────────────────────────┐\n"
        f"  🗓 Sana:   `{dt.day} {month_names[dt.month]} {dt.year}`\n"
        f"  ⏰ Vaqt:   `{dt.strftime('%H : %M')}`\n"
        f"└──────────────────────────────┘\n"
        f"_Quyidagi tugmalar orqali vaqtni sozlang:_"
    )

def build_scheduler_keyboard(lang: str = 'uz') -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("◀️ Kun", callback_data="sched_day_prev"), InlineKeyboardButton("Kun ▶️", callback_data="sched_day_next")],
        [InlineKeyboardButton("➖ 1 soat", callback_data="sched_hour_minus"), InlineKeyboardButton("➕ 1 soat", callback_data="sched_hour_plus")],
        [InlineKeyboardButton("➖ 15 daq", callback_data="sched_min_minus"), InlineKeyboardButton("➕ 15 daq", callback_data="sched_min_plus")],
        [InlineKeyboardButton("⚡ Bugun", callback_data="sched_jump_today"), InlineKeyboardButton("⚡ Ertaga", callback_data="sched_jump_tmrw"), InlineKeyboardButton("⚡ 1 hafta", callback_data="sched_jump_week")],
        [InlineKeyboardButton("🗓 Taqvimdan tanlash", callback_data="sched_open_cal")],
        [InlineKeyboardButton("✅ TASDIQLASH VA SAQLASH", callback_data="sched_confirm")],
        [InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")],
    ])

def build_month_calendar(year: int, month: int, lang: str = 'uz') -> InlineKeyboardMarkup:
    m_names = ["", "Yanvar", "Fevral", "Mart", "Aprel", "May", "Iyun", "Iyul", "Avgust", "Sentyabr", "Oktyabr", "Noyabr", "Dekabr"]
    w_days = ["Du", "Se", "Cho", "Pa", "Ju", "Sha", "Ya"]
    rows = [
        [
            InlineKeyboardButton("◀️", callback_data=f"cal_nav_{year if month > 1 else year - 1}_{month - 1 if month > 1 else 12}"),
            InlineKeyboardButton(f"{m_names[month]} {year}", callback_data="cal_ignore"),
            InlineKeyboardButton("▶️", callback_data=f"cal_nav_{year if month < 12 else year + 1}_{month + 1 if month < 12 else 1}"),
        ],
        [InlineKeyboardButton(d, callback_data="cal_ignore") for d in w_days]
    ]
    cal = calendar.monthcalendar(year, month)
    today = datetime.now()
    for week in cal:
        r = []
        for d in week:
            if d == 0:
                r.append(InlineKeyboardButton(" ", callback_data="cal_ignore"))
            else:
                label = f"•{d}•" if (today.year == year and today.month == month and today.day == d) else str(d)
                r.append(InlineKeyboardButton(label, callback_data=f"cal_pick_{year:04d}-{month:02d}-{d:02d}"))
        rows.append(r)
    rows.append([InlineKeyboardButton("🔙 Orqaga", callback_data="cal_back_panel")])
    return InlineKeyboardMarkup(rows)

async def safe_edit_text_markup(message, text: str, reply_markup=None, parse_mode=None):
    try:
        await message.edit_text(text, reply_markup=reply_markup, parse_mode=parse_mode)
    except Exception:
        try:
            await message.edit_text(text.replace("*", "").replace("_", "").replace("`", ""), reply_markup=reply_markup)
        except Exception:
            pass

# =====================================================================
# MENÜLER VE METİNLER
# =====================================================================
TEXTS = {
    'uz': {
        'welcome': "Assalomu alaykum! Nun Botga xush kelibsiz.\nQuyidagi menyudan kerakli boʻlimni tanlang:",
        'menu_title': "📋 Asosiy menyu:",
        'btn_video': "🎬 Video yuklash",
        'btn_prayer': "🕌 Namoz vaqtlari",
        'btn_pdf_hub': "📄 PDF & Hujjatlar",
        'btn_exam': "🎓 Imtihon & Taymer",
        'btn_schedule_img': "🗓️ Dars Jadvali (Rasm)",
        'btn_pomodoro': "⏱️ Pomodoro & Eslatma",
        'btn_adhkar': "📿 Zikrlar",
        'btn_translit': "🔤 Krill ⇄ Lotin",
        'btn_lang': "🌐 Tilni tanlash",
    },
    'tr': {
        'welcome': "Merhaba! Nun Bota hoş geldiniz.\nAşağıdaki menüden işlem seçiniz:",
        'menu_title': "📋 Ana Menü:",
        'btn_video': "🎬 Video İndir",
        'btn_prayer': "🕌 Namaz Vakitleri",
        'btn_pdf_hub': "📄 PDF & Belge Araçları",
        'btn_exam': "🎓 Sınav & Geri Sayım",
        'btn_schedule_img': "🗓️ Ders Programı (Görsel)",
        'btn_pomodoro': "⏱️ Pomodoro & Hatırlatıcı",
        'btn_adhkar': "📿 Zikirler",
        'btn_translit': "🔤 Kiril ⇄ Latin",
        'btn_lang': "🌐 Dil Seçimi",
    }
}

def get_text(user_id, key, context=None):
    lang = get_user_lang(user_id, context)
    return TEXTS.get(lang, TEXTS['uz']).get(key, '')

def get_reply_menu(user_id, context=None):
    lang = get_user_lang(user_id, context)
    t = TEXTS.get(lang, TEXTS['uz'])
    return ReplyKeyboardMarkup([
        [KeyboardButton(t['btn_video']), KeyboardButton(t['btn_prayer'])],
        [KeyboardButton(t['btn_pdf_hub']), KeyboardButton(t['btn_exam'])],
        [KeyboardButton(t['btn_pomodoro']), KeyboardButton(t['btn_schedule_img'])],
        [KeyboardButton(t['btn_adhkar']), KeyboardButton(t['btn_translit'])],
        [KeyboardButton(t['btn_lang'])]
    ], resize_keyboard=True)

def get_pdf_hub_keyboard(lang: str = 'uz'):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("📸 Rasm / Word / Excel ➔ PDF", callback_data="pdf_act_to_pdf")],
        [InlineKeyboardButton("📑 PDF ➔ Word (.docx)", callback_data="pdf_act_to_word")],
        [InlineKeyboardButton("🗂️ PDF Birlashtirish (Merge)", callback_data="pdf_act_merge")],
        [InlineKeyboardButton("✂️ PDF Ajratish (Split)", callback_data="pdf_act_split")],
        [InlineKeyboardButton("🔍 Rasmdan Matn Olish (OCR)", callback_data="pdf_act_ocr")],
    ])

def get_pomodoro_keyboard():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🍅 25 Dk Ish", callback_data="pomo_25"), InlineKeyboardButton("☕ 5 Dk Mola", callback_data="pomo_5")],
        [InlineKeyboardButton("🍅 50 Dk Ish", callback_data="pomo_50"), InlineKeyboardButton("☕ 10 Dk Mola", callback_data="pomo_10")],
        [InlineKeyboardButton("⏰ Yangi Eslatma Qoʻshish", callback_data="remind_add"), InlineKeyboardButton("📋 Eslatmalarim", callback_data="remind_list")]
    ])

# =====================================================================
# VİDEO İNDİRME MOTORU
# =====================================================================
SUPPORTED_PLATFORMS = [r'(?:instagram\.com)', r'(?:tiktok\.com)', r'(?:facebook\.com|fb\.watch|fb\.gg)', r'(?:twitter\.com|x\.com)']

def is_supported_url(url: str) -> bool:
    return any(re.search(p, url, re.IGNORECASE) for p in SUPPORTED_PLATFORMS)

def download_media_sync(url: str, download_dir: str):
    opts = {
        'outtmpl': os.path.join(download_dir, 'media_%(id)s.%(ext)s'),
        'quiet': True, 'no_warnings': True, 'nocheckcertificate': True,
        'format': 'best[ext=mp4]/bestvideo[ext=mp4]+bestaudio/best',
    }
    if shutil.which('ffmpeg'):
        opts['merge_output_format'] = 'mp4'
    with yt_dlp.YoutubeDL(opts) as ydl:
        info = ydl.extract_info(url, download=True)
        title = info.get('title', 'Video') if info else 'Video'
        for f in os.listdir(download_dir):
            p = os.path.join(download_dir, f)
            if os.path.isfile(p) and os.path.getsize(p) > 10 * 1024:
                return p, title
    raise RuntimeError("Dosya indirilemedi.")

# =====================================================================
# TELEGRAM HANDLERS
# =====================================================================
async def start_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    cleanup_user_temp_files(context, user_id)
    if str(user_id) not in USER_LANGS:
        save_user_lang(user_id, 'uz')
    await update.message.reply_text(get_text(user_id, 'welcome', context), reply_markup=get_reply_menu(user_id, context))

async def menu_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    cleanup_user_temp_files(context, user_id)
    await update.message.reply_text(get_text(user_id, 'menu_title', context), reply_markup=get_reply_menu(user_id, context))

async def cancel_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    cleanup_user_temp_files(context, user_id)
    await update.message.reply_text("✅ Joriy amal bekor qilindi.", reply_markup=get_reply_menu(user_id, context))

async def handle_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    user_id = query.from_user.id
    data = query.data
    user_lang = get_user_lang(user_id, context)

    # İPTAL VE TEMİZLİK
    if data == "cancel_action":
        cleanup_user_temp_files(context, user_id)
        try: await query.message.delete()
        except Exception: pass
        await context.bot.send_message(chat_id=user_id, text="✅ Amal bekor qilindi.", reply_markup=get_reply_menu(user_id, context))
        return

    # DİL DEĞİŞTİRME
    if data.startswith("lang_"):
        l_code = data.replace("lang_", "", 1).strip()
        save_user_lang(user_id, l_code)
        try: await query.message.delete()
        except Exception: pass
        await context.bot.send_message(chat_id=user_id, text=f"✅ {get_text(user_id, 'welcome', context)}", reply_markup=get_reply_menu(user_id, context))
        return

    # PDF ARAÇLARI SEÇİMLERİ (KESİN MOD KİLİTLENİR)
    if data == "pdf_act_to_pdf":
        cleanup_user_temp_files(context, user_id)
        context.user_data['mode'] = 'convert_to_pdf'
        await query.message.reply_text("📸 *PDF'GA OʻGIRISH REJIMI FAOL*\n\nPDF formatiga oʻtkazmoqchi boʻlgan faylni yuboring (Rasm, Word, Excel yoki TXT):", parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))
        return
    if data == "pdf_act_to_word":
        cleanup_user_temp_files(context, user_id)
        context.user_data['mode'] = 'pdf_to_word'
        await query.message.reply_text("📑 *PDF ➔ WORD REJIMI FAOL*\n\nWord (.docx) formatiga oʻtkazmoqchi boʻlgan PDF faylni yuboring:", parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))
        return
    if data == "pdf_act_merge":
        cleanup_user_temp_files(context, user_id)
        context.user_data['mode'] = 'pdf_merge'
        context.user_data['merge_files'] = []
        await query.message.reply_text("🗂️ *PDF BIRLASHTIRISH REJIMI FAOL*\n\nBirlashtirmoqchi boʻlgan PDF fayllarni birin-ketin yuboring (kamida 2 ta).\nTugatgach, quyidagi tugmani bosing:", parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))
        return
    if data == "pdf_act_split":
        cleanup_user_temp_files(context, user_id)
        context.user_data['mode'] = 'pdf_split_wait_file'
        await query.message.reply_text("✂️ *PDF AJRATISH REJIMI FAOL*\n\nSahifalarini ajratmoqchi boʻlgan PDF faylni yuboring:", parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))
        return
    if data == "pdf_act_ocr":
        cleanup_user_temp_files(context, user_id)
        context.user_data['mode'] = 'ocr'
        await query.message.reply_text("🔍 *RASMDAN MATN CHIQARISH (OCR) FAOL*\n\nMatnini oʻqib olmoqchi boʻlgan rasm yoki hujjatni yuboring:", parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))
        return

    # PDF MERGE İŞLEMİ ÇALIŞTIRMA
    if data == "do_pdf_merge":
        files = context.user_data.get('merge_files', [])
        if len(files) < 2:
            await query.message.reply_text("⚠️ Kamida 2 ta PDF fayl yuborishingiz kerak.")
            return
        status = await query.message.reply_text("⚙️ PDF hujjatlari birlashtirilmoqda...")
        with tempfile.TemporaryDirectory() as tmp_dir:
            out_pdf = os.path.join(tmp_dir, "nun_birlashtirilgan.pdf")
            if merge_pdfs(files, out_pdf):
                with open(out_pdf, "rb") as f:
                    await query.message.reply_document(document=f, filename="nun_birlashtirilgan.pdf", caption="✅ PDF muvaffaqiyatli birlashtirildi!")
            else:
                await query.message.reply_text("❌ PDF fayllarini birlashtirishda xatolik yuz berdi.")
        cleanup_user_temp_files(context, user_id)
        try: await status.delete()
        except Exception: pass
        return

    # DOĞRUDAN YÜKLENEN GÖRSEL İŞLEM SEÇİMİ
    if data == "direct_img_pdf":
        img_p = context.user_data.get('direct_file_path')
        if img_p and os.path.exists(img_p):
            with tempfile.TemporaryDirectory() as tmp_dir:
                out_p = os.path.join(tmp_dir, "nun_rasm.pdf")
                with Image.open(img_p) as im:
                    if im.mode in ("RGBA", "P"): im = im.convert("RGB")
                    im.save(out_p, format="PDF")
                with open(out_p, "rb") as f:
                    await query.message.reply_document(document=f, filename="nun_rasm.pdf", caption="✅ Rasm PDF formatiga oʻtkazildi!")
        cleanup_user_temp_files(context, user_id)
        return

    if data == "direct_img_ocr":
        img_p = context.user_data.get('direct_file_path')
        if img_p and os.path.exists(img_p):
            txt = await asyncio.to_thread(extract_text_from_image, img_p)
            if txt:
                await query.message.reply_text(f"🔍 *RASMDAN OʻQIB OLINGAN MATN:*\n\n`{txt}`", parse_mode="Markdown")
            else:
                await query.message.reply_text("❌ Rasmdan tushunarli matn topilmadi.")
        cleanup_user_temp_files(context, user_id)
        return

    # DOĞRUDAN YÜKLENEN PDF İŞLEM SEÇİMİ
    if data == "direct_pdf_word":
        pdf_p = context.user_data.get('direct_file_path')
        if pdf_p and os.path.exists(pdf_p):
            with tempfile.TemporaryDirectory() as tmp_dir:
                out_docx = os.path.join(tmp_dir, "nun_hujjat.docx")
                if await asyncio.to_thread(pdf_to_docx, pdf_p, out_docx):
                    with open(out_docx, "rb") as f:
                        await query.message.reply_document(document=f, filename="nun_hujjat.docx", caption="✅ PDF muvaffaqiyatli Word formatiga oʻgirildi!")
                else:
                    await query.message.reply_text("❌ PDF ichidan metn topilmadi (skanerlangan boʻlishi mumkin).")
        cleanup_user_temp_files(context, user_id)
        return

    if data == "direct_pdf_split":
        pdf_p = context.user_data.get('direct_file_path')
        if pdf_p and os.path.exists(pdf_p):
            reader = PdfReader(pdf_p)
            p_cnt = len(reader.pages)
            context.user_data['split_file_path'] = pdf_p
            context.user_data['split_max_pages'] = p_cnt
            context.user_data['mode'] = 'pdf_split_wait_range'
            context.user_data.pop('direct_file_path', None)
            await query.message.reply_text(f"📄 PDF qabul qilindi! Jami *{p_cnt}* ta sahifa bor.\n\nQaysi sahifalarni ajratmoqchisiz?\n_(Masalan: `1-3` yoki `2, 4`):_", parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))
            return

    # POMODORO
    if data.startswith("pomo_"):
        mins = int(data.replace("pomo_", "", 1))
        is_break = mins in (5, 10)
        end_time = datetime.now() + timedelta(minutes=mins)
        label = "Dam olish" if is_break else "Dars qilish"
        await query.message.reply_text(f"🍅 *POMODORO BOSHLANDI*\n\n📌 *Rejim:* {label}\n⏳ *Vaqt:* `{mins} daqiqa`\n🏁 *Tugash:* `{end_time.strftime('%H:%M')}`", parse_mode="Markdown")
        asyncio.create_task(pomodoro_timer_task(context.bot, user_id, mins, is_break, user_lang))
        return

    if data == "remind_add":
        cleanup_user_temp_files(context, user_id)
        context.user_data['mode'] = 'remind_input'
        await query.message.reply_text("⏰ Eslatmani quyidagicha yuboring:\n`Kitob o'qish - 18:30` yoki `Dars - 30 daqiqa`", parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))
        return

    if data == "remind_list":
        rems = get_user_reminders(user_id)
        if not rems:
            await query.message.reply_text("Sizda faol eslatmalar yoʻq.")
            return
        lines = ["📋 *FAOL ESLATMALAR:*\n"]
        btns = []
        for r in rems:
            lines.append(f"▫️ {r.get('text')} — `{r.get('time')}`")
            btns.append([InlineKeyboardButton(f"❌ Oʻchirish: {r.get('text')[:15]}", callback_data=f"del_rem_{r.get('id')}")])
        await query.message.reply_text("\n".join(lines), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(btns))
        return

    if data.startswith("del_rem_"):
        r_id = data.replace("del_rem_", "", 1)
        delete_user_reminder(user_id, r_id)
        try: await query.message.delete()
        except Exception: pass
        await query.message.reply_text("🗑️ Eslatma oʻchirildi.")
        return

    # SINAV VE TAKVİM İŞLEMLERİ (SEÇENEK A)
    if data == "exam_add":
        cleanup_user_temp_files(context, user_id)
        context.user_data['mode'] = 'exam_title_input'
        await query.message.reply_text("🎓 *IMTIHON QOʻSHISH*\n\n✍️ Imtihon yoki fanning nomini yozib yuboring:\n_(Masalan: *Oliy Matematika*, *Fizika Final*)_", parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))
        return

    if data.startswith("sched_"):
        cur_dt = context.user_data.get('exam_draft_dt') or ((datetime.now() + timedelta(days=1)).replace(hour=10, minute=0, second=0, microsecond=0))
        title = context.user_data.get('exam_draft_title', 'Imtihon')
        if data == "sched_day_prev": cur_dt -= timedelta(days=1)
        elif data == "sched_day_next": cur_dt += timedelta(days=1)
        elif data == "sched_hour_minus": cur_dt -= timedelta(hours=1)
        elif data == "sched_hour_plus": cur_dt += timedelta(hours=1)
        elif data == "sched_min_minus": cur_dt -= timedelta(minutes=15)
        elif data == "sched_min_plus": cur_dt += timedelta(minutes=15)
        elif data == "sched_jump_today": now = datetime.now(); cur_dt = cur_dt.replace(year=now.year, month=now.month, day=now.day)
        elif data == "sched_jump_tmrw": tmrw = datetime.now() + timedelta(days=1); cur_dt = cur_dt.replace(year=tmrw.year, month=tmrw.month, day=tmrw.day)
        elif data == "sched_jump_week": nxt = datetime.now() + timedelta(days=7); cur_dt = cur_dt.replace(year=nxt.year, month=nxt.month, day=nxt.day)
        elif data == "sched_open_cal":
            await safe_edit_text_markup(query.message, f"🗓 *TAQVIM*\n📌 Fan: *{title}*", build_month_calendar(cur_dt.year, cur_dt.month, user_lang), parse_mode="Markdown")
            return
        elif data == "sched_confirm":
            full_dt = cur_dt.strftime("%Y-%m-%d %H:%M")
            exam_id = add_user_exam(user_id, title, full_dt)
            cleanup_user_temp_files(context, user_id)
            try: await query.message.delete()
            except Exception: pass
            card = format_exam_countdown({'id': exam_id, 'title': title, 'date': full_dt}, user_lang)
            await context.bot.send_message(chat_id=user_id, text=f"✅ *Imtihon saqlandi!*\n\n{card}", parse_mode="Markdown")
            return

        context.user_data['exam_draft_dt'] = cur_dt
        await safe_edit_text_markup(query.message, format_scheduler_card(title, cur_dt, user_lang), build_scheduler_keyboard(user_lang), parse_mode="Markdown")
        return

    if data.startswith("cal_pick_"):
        d_str = data.replace("cal_pick_", "", 1).strip()
        y, m, d = [int(x) for x in d_str.split("-")]
        cur_dt = context.user_data.get('exam_draft_dt') or datetime.now().replace(hour=10, minute=0)
        new_dt = cur_dt.replace(year=y, month=m, day=d)
        context.user_data['exam_draft_dt'] = new_dt
        title = context.user_data.get('exam_draft_title', 'Imtihon')
        await safe_edit_text_markup(query.message, format_scheduler_card(title, new_dt, user_lang), build_scheduler_keyboard(user_lang), parse_mode="Markdown")
        return

    if data == "cal_back_panel":
        cur_dt = context.user_data.get('exam_draft_dt') or datetime.now()
        title = context.user_data.get('exam_draft_title', 'Imtihon')
        await safe_edit_text_markup(query.message, format_scheduler_card(title, cur_dt, user_lang), build_scheduler_keyboard(user_lang), parse_mode="Markdown")
        return

    if data.startswith("cal_nav_"):
        parts = data.split("_")
        _, _, ny, nm = parts
        title = context.user_data.get('exam_draft_title', 'Imtihon')
        await safe_edit_text_markup(query.message, f"🗓 *TAQVIM*\n📌 Fan: *{title}*", build_month_calendar(int(ny), int(nm), user_lang), parse_mode="Markdown")
        return

# =====================================================================
# DOSYA & FOTOĞRAF İŞLEYİCİSİ (KESİN MOD KİLİDİ)
# =====================================================================
async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    mode = context.user_data.get('mode', 'auto')
    doc = update.message.document
    if not doc:
        return

    fname = doc.file_name or "fayl"
    ext = os.path.splitext(fname).lower()
    status = await update.message.reply_text("⚙️ Fayl qabul qilindi...")

    try:
        file_obj = await doc.get_file()
        perm_tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
        perm_tmp.close()
        await file_obj.download_to_drive(perm_tmp.name)

        # 1. OCR MODUNDA BELGE OLARAK GÖNDERİLEN GÖRSELLER
        if mode == 'ocr' and ext in (".jpg", ".jpeg", ".png", ".webp", ".bmp"):
            txt = await asyncio.to_thread(extract_text_from_image, perm_tmp.name)
            if txt:
                await update.message.reply_text(f"🔍 *RASMDAN OʻQIB OLINGAN MATN:*\n\n`{txt}`", parse_mode="Markdown")
            else:
                await update.message.reply_text("❌ Rasmdan tushunarli matn topilmadi.")
            cleanup_user_temp_files(context, user_id)
            try: os.remove(perm_tmp.name)
            except Exception: pass
            return

        # 2. PDF MERGE MODUNDA GÖNDERİLEN DOSYA
        if mode == 'pdf_merge' and ext == ".pdf":
            if 'merge_files' not in context.user_data:
                context.user_data['merge_files'] = []
            context.user_data['merge_files'].append(perm_tmp.name)
            cnt = len(context.user_data['merge_files'])
            await update.message.reply_text(
                f"📥 *{cnt}-PDF fayli qabul qilindi!*\n_(Nomi: {fname})_\n\nYana PDF yuborishingiz mumkin yoki birlashtirishni boshlang:",
                parse_mode="Markdown",
                reply_markup=InlineKeyboardMarkup([
                    [InlineKeyboardButton(f"✅ Birlashtirish ({cnt} ta PDF)", callback_data="do_pdf_merge")],
                    [InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]
                ])
            )
            return

        # 3. PDF SPLIT MODUNDA GÖNDERİLEN DOSYA
        if mode == 'pdf_split_wait_file' and ext == ".pdf":
            reader = PdfReader(perm_tmp.name)
            p_cnt = len(reader.pages)
            context.user_data['split_file_path'] = perm_tmp.name
            context.user_data['split_max_pages'] = p_cnt
            context.user_data['mode'] = 'pdf_split_wait_range'
            await update.message.reply_text(
                f"📄 *PDF qabul qilindi!*\nJami sahifalar soni: *{p_cnt}* ta.\n\nQaysi sahifalarni ajratmoqchisiz?\nMasalan:\n▫️ `1-3`\n▫️ `2, 4`\n▫️ `1`",
                parse_mode="Markdown",
                reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]])
            )
            return

        # 4. PDF TO WORD MODUNDA GÖNDERİLEN DOSYA
        if mode == 'pdf_to_word' and ext == ".pdf":
            with tempfile.TemporaryDirectory() as tmp_dir:
                out_docx = os.path.join(tmp_dir, "nun_hujjat.docx")
                if await asyncio.to_thread(pdf_to_docx, perm_tmp.name, out_docx):
                    with open(out_docx, "rb") as f:
                        await update.message.reply_document(document=f, filename="nun_hujjat.docx", caption="✅ PDF Word formatiga oʻtkazildi!")
                else:
                    await update.message.reply_text("❌ PDF ichida matn topilmadi.")
            cleanup_user_temp_files(context, user_id)
            try: os.remove(perm_tmp.name)
            except Exception: pass
            return

        # 5. CONVERT TO PDF MODUNDA GELEN DOSYALAR
        if mode == 'convert_to_pdf':
            with tempfile.TemporaryDirectory() as tmp_dir:
                out_pdf = os.path.join(tmp_dir, "nun_hujjat.pdf")
                ok = False
                if ext in (".docx", ".doc"): ok = await asyncio.to_thread(docx_to_pdf, perm_tmp.name, out_pdf)
                elif ext in (".xlsx", ".xls"): ok = await asyncio.to_thread(xlsx_to_pdf, perm_tmp.name, out_pdf)
                elif ext == ".txt": ok = await asyncio.to_thread(txt_to_pdf, perm_tmp.name, out_pdf)
                elif ext in (".jpg", ".jpeg", ".png"):
                    with Image.open(perm_tmp.name) as im:
                        if im.mode in ("RGBA", "P"): im = im.convert("RGB")
                        im.save(out_pdf, format="PDF")
                    ok = True
                if ok and os.path.exists(out_pdf):
                    with open(out_pdf, "rb") as f:
                        await update.message.reply_document(document=f, filename="nun_hujjat.pdf", caption="✅ PDF hujjati muvaffaqiyatli tayyorlandi!")
                else:
                    await update.message.reply_text("❌ Faylni PDF ga oʻgirishda xatolik yuz berdi.")
            cleanup_user_temp_files(context, user_id)
            try: os.remove(perm_tmp.name)
            except Exception: pass
            return

        # 6. HİÇBİR MOD SEÇİLMEDEN DOĞRUDAN ATILAN BELGELER
        if ext == ".pdf":
            context.user_data['direct_file_path'] = perm_tmp.name
            await update.message.reply_text(
                f"📑 *PDF hujjati qabul qilindi* (`{fname}`).\nNima qilmoqchisiz?",
                parse_mode="Markdown",
                reply_markup=InlineKeyboardMarkup([
                    [InlineKeyboardButton("📄 Word (.docx) ga oʻgirish", callback_data="direct_pdf_word")],
                    [InlineKeyboardButton("✂️ Sahifalarni ajratish (Split)", callback_data="direct_pdf_split")],
                    [InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]
                ])
            )
            return

        if ext in (".jpg", ".jpeg", ".png", ".webp"):
            context.user_data['direct_file_path'] = perm_tmp.name
            await update.message.reply_text(
                f"📸 *Rasm qabul qilindi* (`{fname}`).\nQaysi amalni bajarmoqchisiz?",
                parse_mode="Markdown",
                reply_markup=InlineKeyboardMarkup([
                    [InlineKeyboardButton("📄 PDF ga aylantirish", callback_data="direct_img_pdf")],
                    [InlineKeyboardButton("🔍 Matnni oʻqish (OCR)", callback_data="direct_img_ocr")],
                    [InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]
                ])
            )
            return

        # Word, Excel, TXT doğrudan PDF'e çevrilir
        if ext in (".docx", ".xlsx", ".txt"):
            with tempfile.TemporaryDirectory() as tmp_dir:
                out_pdf = os.path.join(tmp_dir, "nun_hujjat.pdf")
                ok = False
                if ext == ".docx": ok = await asyncio.to_thread(docx_to_pdf, perm_tmp.name, out_pdf)
                elif ext == ".xlsx": ok = await asyncio.to_thread(xlsx_to_pdf, perm_tmp.name, out_pdf)
                elif ext == ".txt": ok = await asyncio.to_thread(txt_to_pdf, perm_tmp.name, out_pdf)
                if ok and os.path.exists(out_pdf):
                    with open(out_pdf, "rb") as f:
                        await update.message.reply_document(document=f, filename="nun_hujjat.pdf", caption="✅ PDF hujjati tayyorlandi!")
            try: os.remove(perm_tmp.name)
            except Exception: pass
            return

        try: os.remove(perm_tmp.name)
        except Exception: pass
        await update.message.reply_text("⚠️ Nomaʼlum fayl formati.")

    except Exception as e:
        print(f"Hujjat xatosi: {e}")
        await update.message.reply_text("❌ Faylni qayta ishlashda xatolik yuz berdi.")
    finally:
        try: await status.delete()
        except Exception: pass

async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    mode = context.user_data.get('mode', 'auto')
    photo = update.message.photo[-1]
    status = await update.message.reply_text("⚙️ Rasm qabul qilindi...")

    try:
        file_obj = await photo.get_file()
        perm_tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".jpg")
        perm_tmp.close()
        await file_obj.download_to_drive(perm_tmp.name)

        # 1. OCR MODU AKTİFSE DOĞRUDAN METNİ OKU
        if mode == 'ocr':
            txt = await asyncio.to_thread(extract_text_from_image, perm_tmp.name)
            if txt:
                await update.message.reply_text(f"🔍 *RASMDAN OʻQIB OLINGAN MATN:*\n\n`{txt}`", parse_mode="Markdown")
            else:
                await update.message.reply_text("❌ Rasmdan tushunarli matn topilmadi.")
            cleanup_user_temp_files(context, user_id)
            try: os.remove(perm_tmp.name)
            except Exception: pass
            return

        # 2. PDF DÖNÜŞTÜRME MODU AKTİFSE
        if mode == 'convert_to_pdf':
            with tempfile.TemporaryDirectory() as tmp_dir:
                out_pdf = os.path.join(tmp_dir, "nun_rasm.pdf")
                with Image.open(perm_tmp.name) as im:
                    if im.mode in ("RGBA", "P"): im = im.convert("RGB")
                    im.save(out_pdf, format="PDF")
                with open(out_pdf, "rb") as f:
                    await update.message.reply_document(document=f, filename="nun_rasm.pdf", caption="✅ Rasm PDF formatiga oʻtkazildi!")
            cleanup_user_temp_files(context, user_id)
            try: os.remove(perm_tmp.name)
            except Exception: pass
            return

        # 3. MOD SEÇİLMEDEN ATILDIYSA SOR
        context.user_data['direct_file_path'] = perm_tmp.name
        await update.message.reply_text(
            "📸 *Rasm qabul qilindi.*\nQaysi amalni bajarmoqchisiz?",
            parse_mode="Markdown",
            reply_markup=InlineKeyboardMarkup([
                [InlineKeyboardButton("📄 PDF ga aylantirish", callback_data="direct_img_pdf")],
                [InlineKeyboardButton("🔍 Matnni oʻqish (OCR)", callback_data="direct_img_ocr")],
                [InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]
            ])
        )

    except Exception as e:
        print(f"Rasm xatosi: {e}")
        await update.message.reply_text("❌ Rasmni qayta ishlashda xatolik yuz berdi.")
    finally:
        try: await status.delete()
        except Exception: pass

# =====================================================================
# METİN MESAJ YÖNLENDİRİCİSİ (KATI KORUMALI)
# =====================================================================
async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.message or not update.message.text:
        return

    user_id = update.effective_user.id
    chat_id = update.message.chat_id
    raw_text = update.message.text.strip()
    user_lang = get_user_lang(user_id, context)

    # 1. Menü Butonları (Her tıklandığında önceki geçici modları temizler)
    btn_actions = {
        'btn_video': lambda: update.message.reply_text("🔗 Instagram, TikTok, Facebook yoki X (Twitter) havolasini yuboring:"),
        'btn_prayer': lambda: set_mode_and_reply(context, 'prayer', update, "🕌 Namoz vaqtlarini bilmoqchi boʻlgan shahar nomini yozing:"),
        'btn_pdf_hub': lambda: update.message.reply_text("📄 *NUN PROJECT // PDF & HUJJATLAR MARKAZI*\n\nAmalni tanlang:", parse_mode="Markdown", reply_markup=get_pdf_hub_keyboard(user_lang)),
        'btn_exam': lambda: show_exams_ui(user_id, user_lang, update),
        'btn_schedule_img': lambda: set_mode_and_reply(context, 'schedule_img_input', update, "🗓️ *DARS JADVALI GÖRSELİ*\n\nDars jadvalingizni kunlar boʻyicha yozib yuboring (Masalan: Dushanba: 09:00 Matematika...):\nBot 1080x1920 kilit ekrani formatiga aylantiradi.", parse_mode="Markdown"),
        'btn_pomodoro': lambda: update.message.reply_text("⏱️ *POMODORO & ESLATMA MARKAZI*", parse_mode="Markdown", reply_markup=get_pomodoro_keyboard()),
        'btn_adhkar': lambda: update.message.reply_text("📿 Zikrlarni tanlang:", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("🌅 Tonggi zikrlar", callback_data="adhkar_morning"), InlineKeyboardButton("🌇 Kechki zikrlar", callback_data="adhkar_evening")]])),
        'btn_translit': lambda: set_mode_and_reply(context, 'translit', update, "✍️ Matningizni yuboring, avtomatik Kirill ⇄ Lotin oʻgirib beraman:"),
        'btn_lang': lambda: update.message.reply_text("Tilni tanlang:", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("🇺🇿 O'zbekcha", callback_data="lang_uz"), InlineKeyboardButton("🇷🇺 Русский", callback_data="lang_ru")], [InlineKeyboardButton("🇹🇷 Türkçe", callback_data="lang_tr"), InlineKeyboardButton("🇬🇧 English", callback_data="lang_en")]]))
    }

    for b_key, action_func in btn_actions.items():
        if raw_text in [TEXTS[l].get(b_key, '') for l in TEXTS]:
            cleanup_user_temp_files(context, user_id)
            await action_func()
            return

    # 2. Medya İndirme Linki
    if is_supported_url(raw_text):
        url = re.search(r'https?://[^\s]+', raw_text).group(0)
        status = await update.message.reply_text("⏳ Video yuklab olinmoqda...")
        try:
            with tempfile.TemporaryDirectory() as tmp_dir:
                f_path, title = await asyncio.to_thread(download_media_sync, url, tmp_dir)
                clean_t = re.sub(r'[\\/*?:"<>|]', '', title)[:60]
                with open(f_path, 'rb') as f:
                    await context.bot.send_video(chat_id=chat_id, video=f, caption=f"🎬 {clean_t}", supports_streaming=True)
            await status.delete()
        except Exception as e:
            print(f"Video xatosi: {e}")
            await status.edit_text("❌ Videoni yuklab olishda xatolik yuz berdi.")
        return

    mode = context.user_data.get('mode', 'auto')

    # 3. KORUMALI AKTİF MODLAR (ASLA BAŞKA YERE SIZAMAZ)
    if mode == 'exam_title_input':
        context.user_data['exam_draft_title'] = raw_text.strip()
        context.user_data['mode'] = 'exam_schedule_panel'
        init_dt = (datetime.now() + timedelta(days=1)).replace(hour=10, minute=0, second=0, microsecond=0)
        context.user_data['exam_draft_dt'] = init_dt
        await update.message.reply_text(format_scheduler_card(raw_text.strip(), init_dt, user_lang), parse_mode="Markdown", reply_markup=build_scheduler_keyboard(user_lang))
        return

    if mode == 'exam_schedule_panel':
        cur_dt = context.user_data.get('exam_draft_dt', datetime.now())
        title = context.user_data.get('exam_draft_title', 'Imtihon')
        await update.message.reply_text(format_scheduler_card(title, cur_dt, user_lang), parse_mode="Markdown", reply_markup=build_scheduler_keyboard(user_lang))
        return

    if mode == 'pdf_split_wait_range':
        file_path = context.user_data.get('split_file_path')
        if file_path and os.path.exists(file_path):
            with tempfile.TemporaryDirectory() as tmp_dir:
                out_pdf = os.path.join(tmp_dir, "nun_ajratilgan.pdf")
                if split_pdf(file_path, raw_text, out_pdf):
                    with open(out_pdf, "rb") as f:
                        await update.message.reply_document(document=f, filename="nun_ajratilgan.pdf", caption=f"✅ `{raw_text}` sahifalari ajratib olindi!", parse_mode="Markdown")
                else:
                    await update.message.reply_text("❌ Sahifa oraligʻi notoʻgʻri. Masalan: `1-3` yoki `2, 4`")
                    return
        cleanup_user_temp_files(context, user_id)
        return

    if mode == 'schedule_img_input':
        status = await update.message.reply_text("⚙️ Kilit ekrani fon rasmi tayyorlanmoqda...")
        parsed_data = parse_schedule_text(raw_text)
        with tempfile.TemporaryDirectory() as tmp_dir:
            out_img = os.path.join(tmp_dir, "timetable_wallpaper.png")
            generate_schedule_wallpaper(parsed_data, out_img, "DARS JADVALI")
            with open(out_img, "rb") as f_photo, open(out_img, "rb") as f_doc:
                await update.message.reply_photo(photo=f_photo, caption="📱 *Dars Jadvali (Kilit Ekrani)*", parse_mode="Markdown")
                await update.message.reply_document(document=f_doc, filename="dars_jadvali_1080x1920.png", caption="🖼️ Asl sifatdagi rasm")
        cleanup_user_temp_files(context, user_id)
        try: await status.delete()
        except Exception: pass
        return

    if mode == 'remind_input':
        target_dt = None
        rem_text = raw_text
        if "-" in raw_text:
            rem_text, _, time_part = raw_text.partition("-")
            rem_text = rem_text.strip()
            time_part = time_part.strip()
            m_min = re.search(r"(\d+)\s*(?:daqiqa|dakika|min|m)", time_part, re.IGNORECASE)
            if m_min: target_dt = datetime.now() + timedelta(minutes=int(m_min.group(1)))
            else:
                m_t = re.search(r"(\d{1,2})[:.](\d{2})", time_part)
                if m_t:
                    target_dt = datetime.now().replace(hour=int(m_t.group(1)), minute=int(m_t.group(2)), second=0)
                    if target_dt < datetime.now(): target_dt += timedelta(days=1)

        if not target_dt:
            m_min = re.search(r"(\d+)\s*(?:daqiqa|dakika|min|m)", raw_text, re.IGNORECASE)
            if m_min:
                target_dt = datetime.now() + timedelta(minutes=int(m_min.group(1)))
                rem_text = re.sub(r"(\d+)\s*(?:daqiqa|dakika|min|m)", "", raw_text, flags=re.IGNORECASE).strip()

        if target_dt:
            dt_str = target_dt.strftime("%Y-%m-%d %H:%M")
            add_user_reminder(user_id, rem_text, dt_str)
            cleanup_user_temp_files(context, user_id)
            await update.message.reply_text(f"✅ *Eslatma oʻrnatildi!*\n\n📌 *Vazifa:* {rem_text}\n⏰ *Vaqt:* `{dt_str}`", parse_mode="Markdown")
        else:
            await update.message.reply_text("⚠️ Vaqtni aniqlab boʻlmadi. Masalan:\n`Kitob o'qish - 18:30` yoki `Dars - 30 daqiqa`", parse_mode="Markdown")
        return

    if mode == 'prayer':
        timings, d_name, dt_s, h_s, src = await fetch_prayer_times(raw_text)
        if timings:
            await update.message.reply_text(format_prayer_card(d_name, timings, dt_s, h_s, src, user_lang), parse_mode="Markdown")
            cleanup_user_temp_files(context, user_id)
            return
        await update.message.reply_text("❌ Shahar topilmadi. Shahar nomini toʻgʻri yozing.")
        return

    if mode == 'translit' or len(raw_text.split()) >= 3:
        if is_mostly_cyrillic(raw_text):
            await update.message.reply_text(f"🔤 *Lotin:*\n\n{cyrillic_to_latin(raw_text)}", parse_mode="Markdown")
        else:
            await update.message.reply_text(f"🔤 *Кирилл:*\n\n{latin_to_cyrillic(raw_text)}", parse_mode="Markdown")
        return

    # Tanınmayan metinlerde menüyü göster
    await update.message.reply_text(get_text(user_id, 'menu_title', context), reply_markup=get_reply_menu(user_id, context))

async def set_mode_and_reply(context, mode_name, update, text, parse_mode=None):
    context.user_data['mode'] = mode_name
    await update.message.reply_text(text, parse_mode=parse_mode, reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="cancel_action")]]))

async def show_exams_ui(user_id, user_lang, update):
    exams = get_user_exams(user_id)
    cards = [format_exam_countdown(e, user_lang) for e in exams] if exams else ["Sizda hali saqlangan imtihon yoʻq."]
    await update.message.reply_text("🎓 *IMTIHONLAR VA TAYMER*\n\n" + "\n".join(cards), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("➕ Imtihon qoʻshish", callback_data="exam_add")]]))

def main():
    token = os.environ.get("BOT_TOKEN")
    if not token: raise ValueError("BOT_TOKEN ortam değişkeni eksik!")
    load_databases()

    threading.Thread(target=run_health_server, daemon=True).start()
    threading.Thread(target=run_keep_alive_pinger, daemon=True).start()

    app = ApplicationBuilder().token(token).build()
    app.add_handler(CommandHandler("start", start_command))
    app.add_handler(CommandHandler("menu", menu_command))
    app.add_handler(CommandHandler("cancel", cancel_command))
    app.add_handler(CallbackQueryHandler(handle_callback))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))

    loop = asyncio.get_event_loop()
    loop.create_task(reminders_worker(app))

    print("Nun Bot korumalı durum mimarisi ve tam PDF paketiyle devrede!")
    app.run_polling(drop_pending_updates=True)

if __name__ == "__main__":
    main()
